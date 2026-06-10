#!/usr/bin/env python3
"""Vytvoří reformatovaný DOCX z DOC dokumentu Spice must flow."""

import re, os, glob
from io import BytesIO
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY      = RGBColor(0x1A, 0x1A, 0x5C)
DARK_GRAY = RGBColor(0x44, 0x44, 0x44)
LIGHT_GRAY= RGBColor(0x88, 0x88, 0x88)

with open('/tmp/spice_text.txt', encoding='utf-8') as f:
    raw = f.read()

images = sorted(glob.glob('/tmp/spice_ordered/*.jpg') + glob.glob('/tmp/spice_ordered/*.png'),
                key=lambda p: int(os.path.basename(p).split('.')[0]))

segments = raw.split('[pic]')

CHAPTER_RE = re.compile(
    r'^(Spice must flow|Jaderná zápalka|Zuřivost|Noc je kratší než den|'
    r'Bude ti dáno znamení|Tento provázaný.*svět|Vládce, zrozený vládnout světu|'
    r'Kanadští sloni a australská želva|Vrata tlamy jeho).*\(NUC\d+\)$'
)
URL_RE     = re.compile(r'^https?://')
QUOTE_RE   = re.compile(r'^[—–]\s*(Princess|Baron|Irulan|Harkonnen)')
CAPTION_RE = re.compile(r'^(Náhled|Vypadá|Dříve|Symbol|Zde je|Na snímku|'
                          r'Toto je|Takový|Schéma|Zde jsou|A vůbec)')

def add_page_break(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r._r.append(br)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)

def add_image(doc, path):
    try:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after  = Pt(6)
        p.add_run().add_picture(path, width=Cm(14))
    except Exception as e:
        print(f'  Varování: {path}: {e}')

def body_para(doc, text, first=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(5)
    p.paragraph_format.first_line_indent = Pt(0) if first else Cm(1.25)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)
    return p

def parse_segment(seg):
    items = []
    buf = []

    def flush():
        joined = ' '.join(buf).strip()
        if joined:
            items.append(('body', joined))
        buf.clear()

    for line in seg.split('\n'):
        line = line.rstrip()
        if not line:
            flush()
            continue
        if CHAPTER_RE.match(line):
            flush()
            items.append(('chapter', line))
            continue
        if URL_RE.match(line.strip()):
            flush()
            items.append(('url', line.strip()))
            continue
        if QUOTE_RE.match(line):
            flush()
            items.append(('quote_attr', line))
            continue
        if CAPTION_RE.match(line) and len(line) < 100:
            flush()
            items.append(('caption', line))
            continue
        buf.append(line)
    flush()
    return items

# ── Sestavení dokumentu ───────────────────────────────────────────────────────
doc = Document()
section = doc.sections[0]
section.top_margin    = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin   = Cm(3.0)
section.right_margin  = Cm(2.5)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(11.5)

h1 = doc.styles['Heading 1']
h1.font.name = 'Times New Roman'
h1.font.size = Pt(17)
h1.font.bold = True
h1.font.color.rgb = NAVY
h1.paragraph_format.space_before = Pt(0)
h1.paragraph_format.space_after  = Pt(6)

# Titulní strana
tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
tp.paragraph_format.space_before = Pt(100)
tp.paragraph_format.space_after  = Pt(10)
r = tp.add_run('Spice must flow')
r.font.name = 'Times New Roman'; r.font.size = Pt(30)
r.font.bold = True; r.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_before = Pt(0)
sub.paragraph_format.space_after  = Pt(16)
r2 = sub.add_run('He who controls the spice, controls the universe')
r2.font.name = 'Times New Roman'; r2.font.size = Pt(13)
r2.font.italic = True; r2.font.color.rgb = DARK_GRAY

auth = doc.add_paragraph()
auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = auth.add_run('Napsal: crustgroup  ·  Přeložil: Hamilbar')
r3.font.name = 'Times New Roman'; r3.font.size = Pt(11)
r3.font.italic = True; r3.font.color.rgb = DARK_GRAY

add_page_break(doc)

# Úvodní citáty
for line in [l.strip() for l in segments[0].split('\n') if l.strip()]:
    if CHAPTER_RE.match(line):
        break
    if line.startswith('—') or line.startswith('–'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(4)
        r = p.add_run(line)
        r.font.name = 'Times New Roman'; r.font.size = Pt(10)
        r.font.color.rgb = LIGHT_GRAY
    elif not URL_RE.match(line):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(4)
        p.paragraph_format.left_indent  = Cm(2)
        p.paragraph_format.right_indent = Cm(2)
        r = p.add_run(line)
        r.font.name = 'Times New Roman'; r.font.size = Pt(12)
        r.font.italic = True; r.font.color.rgb = DARK_GRAY

img_idx = 0
chapter_count = 0

for seg_idx, seg in enumerate(segments):
    is_first = True
    items = parse_segment(seg)

    for typ, content in items:
        if typ == 'chapter':
            if chapter_count > 0:
                add_page_break(doc)
            chapter_count += 1
            title = re.sub(r'\s*\(NUC\d+\)', '', content).strip()
            num   = re.search(r'NUC(\d+)', content)
            label = f"NUC{num.group(1)}" if num else ''

            h = doc.add_paragraph(style='Heading 1')
            h.clear()
            r = h.add_run(title)
            r.font.name = 'Times New Roman'; r.font.size = Pt(17); r.font.bold = True

            sub_p = doc.add_paragraph()
            sub_p.paragraph_format.space_before = Pt(0)
            sub_p.paragraph_format.space_after  = Pt(12)
            r2 = sub_p.add_run(label)
            r2.font.name = 'Times New Roman'; r2.font.size = Pt(11)
            r2.font.italic = True; r2.font.color.rgb = DARK_GRAY
            is_first = True

        elif typ == 'translator':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after  = Pt(10)
            r = p.add_run(content)
            r.font.name = 'Times New Roman'; r.font.size = Pt(10)
            r.font.italic = True; r.font.color.rgb = DARK_GRAY
            is_first = True

        elif typ == 'url':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after  = Pt(2)
            r = p.add_run(content)
            r.font.name = 'Times New Roman'; r.font.size = Pt(8.5)
            r.font.italic = True; r.font.color.rgb = LIGHT_GRAY

        elif typ in ('quote', 'quote_attr'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after  = Pt(3)
            r = p.add_run(content)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10 if typ == 'quote_attr' else 12)
            r.font.italic = True
            r.font.color.rgb = DARK_GRAY if typ == 'quote' else LIGHT_GRAY

        elif typ == 'caption':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after  = Pt(8)
            r = p.add_run(content)
            r.font.name = 'Times New Roman'; r.font.size = Pt(9.5)
            r.font.italic = True; r.font.color.rgb = DARK_GRAY

        elif typ == 'body':
            body_para(doc, content, first=is_first)
            is_first = False

    # Obrázek za segmentem
    if seg_idx < len(segments) - 1:
        if img_idx < len(images):
            add_image(doc, images[img_idx])
            img_idx += 1

print(f'Použito obrázků: {img_idx}')
out_docx = '/home/user/SoulBrowser/Spice_must_flow_reformatovano.docx'
doc.save(out_docx)
print(f'DOCX uložen: {out_docx}')
